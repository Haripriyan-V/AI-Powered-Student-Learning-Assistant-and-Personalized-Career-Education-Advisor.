"""
Resume Analyzer: In-depth resume parsing, text extraction, skill normalization,
ATS scoring rubric, and target career alignment.
"""
import re
import json
import zipfile
import xml.etree.ElementTree as ET
import logging
from typing import Dict, List, Any, Tuple

from ai.skill_engine import normalize_skill_name, SKILL_NORMALIZATION_MAP

logger = logging.getLogger(__name__)


class ResumeParser:
    """Extracts raw text from PDF, DOCX, and TXT files using zero or minimal external dependencies."""

    @classmethod
    def extract_text(cls, file_obj, filename: str) -> str:
        filename_lower = (filename or '').lower()
        try:
            if hasattr(file_obj, 'seek'):
                file_obj.seek(0)
        except Exception:
            pass

        if filename_lower.endswith('.txt') or filename_lower.endswith('.md'):
            try:
                content = file_obj.read()
                if isinstance(content, bytes):
                    return content.decode('utf-8', errors='ignore')
                return str(content)
            except Exception as e:
                logger.error("Error reading text file: %s", e)
                return ""

        elif filename_lower.endswith('.docx') or filename_lower.endswith('.doc'):
            # First try python-docx if available and valid .docx
            if filename_lower.endswith('.docx'):
                try:
                    import docx
                    if hasattr(file_obj, 'seek'):
                        file_obj.seek(0)
                    doc = docx.Document(file_obj)
                    paragraphs_text = [p.text for p in doc.paragraphs if p.text]
                    for table in doc.tables:
                        for row in table.rows:
                            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                            if row_cells:
                                paragraphs_text.append(" | ".join(row_cells))
                    extracted = "\n".join(paragraphs_text)
                    if extracted.strip():
                        return extracted
                except Exception as docx_err:
                    logger.info("python-docx parsing note: %s. Trying zipfile parser.", docx_err)

            # DOCX is a zip file containing word/document.xml
            try:
                if hasattr(file_obj, 'seek'):
                    file_obj.seek(0)
                with zipfile.ZipFile(file_obj) as z:
                    xml_content = z.read('word/document.xml')
                    tree = ET.fromstring(xml_content)
                    texts = []
                    for elem in tree.iter():
                        if elem.tag.endswith('t'):
                            if elem.text:
                                texts.append(elem.text)
                    extracted = " ".join(texts)
                    if extracted.strip():
                        return extracted
            except Exception as e:
                logger.warning("DOCX extraction via zipfile failed (%s), attempting text stream fallback", e)

            # Fallback for binary .doc or malformed docx
            try:
                if hasattr(file_obj, 'seek'):
                    file_obj.seek(0)
                raw_bytes = file_obj.read()
                if isinstance(raw_bytes, bytes):
                    clean_words = re.findall(r'[A-Za-z0-9\.\,\+\#\-\_\/\@]{2,}', raw_bytes.decode('latin1', errors='ignore'))
                    return " ".join(clean_words[:2000])
            except Exception as e:
                logger.error("Text stream fallback for doc failed: %s", e)
                return ""

        elif filename_lower.endswith('.pdf'):
            # Try pypdf
            try:
                if hasattr(file_obj, 'seek'):
                    file_obj.seek(0)
                import pypdf
                reader = pypdf.PdfReader(file_obj)
                pages = [page.extract_text() or '' for page in reader.pages]
                text = "\n".join(pages)
                if text.strip():
                    return text
            except Exception as e:
                logger.warning("pypdf extraction failed: %s, attempting fallback", e)

            # Fallback for PDF text stream reading
            try:
                if hasattr(file_obj, 'seek'):
                    file_obj.seek(0)
                raw_bytes = file_obj.read()
                # Simple extraction of string objects from PDF streams
                strings = re.findall(rb'\(([^\(\)]*)\)\s*T[jJ]', raw_bytes)
                if strings:
                    extracted = " ".join(s.decode('latin1', errors='ignore') for s in strings)
                    if len(extracted.strip()) > 50:
                        return extracted

                # Fallback utf-8/latin1 decode
                text = raw_bytes.decode('latin1', errors='ignore')
                clean_words = re.findall(r'[A-Za-z0-9\.\,\+\#\-\_\/\@]{2,}', text)
                return " ".join(clean_words[:2000])
            except Exception as e:
                logger.error("Fallback PDF extraction failed: %s", e)
                return ""

        return ""


class ResumeAnalyzer:
    """Evaluates resume text for skills, ATS compatibility, and career readiness."""

    # Common Section Headings
    SECTION_PATTERNS = {
        'education': r'(education|academic|qualification|university|college|degree)',
        'experience': r'(experience|employment|work history|professional experience|internship)',
        'projects': r'(projects|personal projects|academic projects|portfolio)',
        'skills': r'(skills|technical skills|technologies|proficiencies|core competencies)',
        'summary': r'(summary|objective|profile|about me|professional summary)',
        'certifications': r'(certifications|certificates|licenses|courses)',
    }

    # Action verbs indicating measurable impact
    ACTION_VERBS = [
        'developed', 'designed', 'implemented', 'built', 'created', 'optimized',
        'reduced', 'increased', 'improved', 'engineered', 'led', 'architected',
        'automated', 'deployed', 'managed', 'analyzed', 'trained', 'scaled',
        'integrated', 'spearheaded', 'orchestrated'
    ]

    @classmethod
    def analyze_resume(cls, file_obj, filename: str, target_career=None) -> Dict[str, Any]:
        """Runs end-to-end analysis on an uploaded resume."""
        text = ResumeParser.extract_text(file_obj, filename)
        if not text or len(text.strip()) < 20:
            return cls._empty_or_failed_result(filename, "Could not extract readable text from the uploaded resume.")

        text_lower = text.lower()

        # 1. Extract Candidate Name (heuristics)
        extracted_name = cls._extract_candidate_name(text)

        # 2. Extract and Normalize Skills
        extracted_skills, normalized_skills = cls._extract_skills(text_lower)

        # 3. Detect Sections
        sections_found = cls._detect_sections(text_lower)

        # 4. Extract Education History
        education = cls._extract_education(text)

        # 5. Extract Experience / Projects summary
        experience = cls._extract_experience(text)
        projects = cls._extract_projects(text)

        # 6. Try AI Evaluation via Gemini
        ai_data = None
        try:
            from ai.service import ai_service
            if ai_service.api_key and not ai_service.api_key.startswith('your_'):
                prompt = (
                    "Evaluate this resume text. Return ONLY a valid JSON object without markdown formatting:\n"
                    '{"ats_score": 82, "summary": "Short 2-sentence overview.", "strengths": ["s1", "s2"], "weaknesses": ["w1", "w2"], "skills": ["Python", "SQL"], "missing_skills": ["Docker"], "recommendations": ["r1", "r2"]}\n\n'
                    f"Resume:\n{text[:2500]}"
                )
                system_prompt = "You are an ATS resume evaluator. Return valid JSON only, no markdown, no other text."
                raw_response = ai_service._call_gemini(system_prompt, [], prompt)
                clean_json = raw_response.strip()
                if clean_json.startswith('```'):
                    clean_json = re.sub(r'^```(?:json)?\s*', '', clean_json)
                    clean_json = re.sub(r'\s*```$', '', clean_json)
                json_match = re.search(r'\{[\s\S]*\}', clean_json)
                if json_match:
                    ai_data = json.loads(json_match.group(0))
                else:
                    ai_data = json.loads(clean_json)
        except Exception as e:
            logger.warning("AI resume analysis call failed or unavailable: %s", e)
            ai_data = None

        # Merge AI skills if available
        if ai_data and isinstance(ai_data.get('skills'), list):
            for sk in ai_data['skills']:
                if isinstance(sk, str) and sk.strip():
                    norm = normalize_skill_name(sk.strip())
                    if norm not in normalized_skills:
                        normalized_skills.append(norm)
                    if sk.strip() not in extracted_skills:
                        extracted_skills.append(sk.strip())

        # 7. Compute ATS & Breakdown Scores (with deterministic rubric)
        scores, strengths, improvements, missing_keywords = cls._compute_ats_scores(
            text=text,
            text_lower=text_lower,
            sections=sections_found,
            skills=normalized_skills,
            education=education,
            experience=experience,
            projects=projects,
        )

        # If AI provided valid ATS score, use it
        if ai_data and isinstance(ai_data.get('ats_score'), (int, float)):
            ai_ats = int(ai_data['ats_score'])
            if 0 <= ai_ats <= 100:
                scores['ats'] = ai_ats
                scores['overall'] = min(100, max(scores['overall'], ai_ats))
                scores['breakdown']['ats_compatibility'] = ai_ats

        if ai_data and isinstance(ai_data.get('strengths'), list) and ai_data['strengths']:
            strengths = [str(s) for s in ai_data['strengths'][:4]]

        if ai_data and (ai_data.get('weaknesses') or ai_data.get('recommendations')):
            ai_imps = []
            if isinstance(ai_data.get('weaknesses'), list):
                ai_imps.extend([str(w) for w in ai_data['weaknesses'][:2]])
            if isinstance(ai_data.get('recommendations'), list):
                ai_imps.extend([str(r) for r in ai_data['recommendations'][:2]])
            if ai_imps:
                improvements = ai_imps

        if ai_data and isinstance(ai_data.get('missing_skills'), list) and ai_data['missing_skills']:
            missing_keywords = [str(m) for m in ai_data['missing_skills'][:5]]

        # 8. Career Alignment
        career_alignment = cls._evaluate_career_alignment(normalized_skills, target_career)

        return {
            'file_name': filename,
            'extracted_name': extracted_name,
            'overall_score': scores['overall'],
            'ats_score': scores['ats'],
            'breakdown_scores': scores['breakdown'],
            'extracted_skills': extracted_skills,
            'normalized_skills': normalized_skills,
            'extracted_education': education,
            'extracted_experience': experience,
            'extracted_projects': projects,
            'strengths': strengths,
            'improvements': improvements,
            'missing_keywords': missing_keywords,
            'career_alignment': career_alignment,
        }

    @classmethod
    def _extract_candidate_name(cls, text: str) -> str:
        lines = [line.strip() for line in text.split('\n') if line.strip()]
        for line in lines[:5]:
            # Look for lines with 2-4 words that look like a person's name
            if re.match(r'^[A-Z][a-z]+(\s+[A-Z][a-z]+){1,3}$', line):
                return line
        return "Student Candidate"

    @classmethod
    def _extract_skills(cls, text_lower: str) -> Tuple[List[str], List[str]]:
        found_canonical = set()
        found_raw = []

        # Check all entries in normalization map
        for alias, canonical in SKILL_NORMALIZATION_MAP.items():
            pattern = r'(?<![a-zA-Z0-9_\-\.\#])' + re.escape(alias) + r'(?![a-zA-Z0-9_\-\.\#])'
            if re.search(pattern, text_lower):
                found_raw.append(alias)
                found_canonical.add(canonical)

        canonical_list = sorted(list(found_canonical))
        return found_raw[:25], canonical_list

    @classmethod
    def _detect_sections(cls, text_lower: str) -> Dict[str, bool]:
        sections = {}
        for sec, pattern in cls.SECTION_PATTERNS.items():
            sections[sec] = bool(re.search(r'\b' + pattern + r'\b', text_lower))
        return sections

    @classmethod
    def _extract_education(cls, text: str) -> List[Dict[str, str]]:
        results = []
        patterns = [
            r'(Bachelor|Master|B\.?Tech|B\.?E\.?|B\.?Sc|M\.?Sc|B\.?C\.?A|M\.?C\.?A|Diploma|Ph\.?D)[\s\w,.-]{0,80}',
            r'(Computer Science|Information Technology|Data Science|Artificial Intelligence|Mechanical|Electrical|Electronics)[\s\w,.-]{0,60}',
        ]
        for p in patterns:
            matches = re.finditer(p, text, re.IGNORECASE)
            for m in matches:
                val = m.group(0).strip()
                if len(val) > 4 and val not in [r.get('degree') for r in results]:
                    results.append({'degree': val, 'institution': 'Academic Institution'})
                    if len(results) >= 3:
                        break
        return results or [{'degree': 'Degree / Coursework in Technology', 'institution': 'University / College'}]

    @classmethod
    def _extract_experience(cls, text: str) -> List[Dict[str, str]]:
        results = []
        matches = re.finditer(r'(Intern|Developer|Engineer|Assistant|Trainee|Lead|Specialist)[\s\w,.-]{0,50}', text, re.IGNORECASE)
        for m in matches:
            val = m.group(0).strip()
            if len(val) > 5 and val not in [r.get('role') for r in results]:
                results.append({'role': val, 'duration': 'Recent'})
                if len(results) >= 3:
                    break
        return results

    @classmethod
    def _extract_projects(cls, text: str) -> List[Dict[str, str]]:
        results = []
        matches = re.finditer(r'(Project|System|Application|Portal|Platform|Model|App)[\s\w,.-]{0,40}', text, re.IGNORECASE)
        for m in matches:
            val = m.group(0).strip()
            if len(val) > 6 and val not in [r.get('title') for r in results]:
                results.append({'title': val})
                if len(results) >= 4:
                    break
        return results

    @classmethod
    def _compute_ats_scores(cls, text: str, text_lower: str, sections: dict, skills: list, education: list, experience: list, projects: list):
        # 1. Contact Information (5% of total score)
        has_email = bool(re.search(r'[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+', text))
        has_phone = bool(re.search(r'(\+?\d{1,3}[-.\s]?)?\(?\d{2,4}\)?[-.\s]?\d{3,4}[-.\s]?\d{3,4}', text))
        has_link = bool(re.search(r'(linkedin\.com|github\.com|portfolio|http[s]?://)', text_lower))
        contact_pts = 0
        if has_email:
            contact_pts += 2
        if has_phone:
            contact_pts += 2
        if has_link:
            contact_pts += 1
        contact_pct = min(100, int((contact_pts / 5.0) * 100))

        # 2. Skills (20% of total score)
        skill_count = len(skills)
        if skill_count >= 8:
            skill_score = 20
        elif skill_count >= 5:
            skill_score = 16
        elif skill_count >= 3:
            skill_score = 12
        elif skill_count >= 1:
            skill_score = 8
        else:
            skill_score = 4
        skill_pct = min(100, int((skill_score / 20.0) * 100))

        # 3. Experience (15% of total score)
        exp_count = len(experience) or (1 if sections.get('experience') else 0)
        if exp_count >= 2:
            experience_score = 15
        elif exp_count == 1:
            experience_score = 11
        elif sections.get('experience'):
            experience_score = 8
        else:
            experience_score = 4
        experience_pct = min(100, int((experience_score / 15.0) * 100))

        # 4. Projects (15% of total score)
        proj_count = len(projects) or (1 if sections.get('projects') else 0)
        if proj_count >= 2:
            project_score = 15
        elif proj_count == 1:
            project_score = 11
        elif sections.get('projects'):
            project_score = 8
        else:
            project_score = 4
        project_pct = min(100, int((project_score / 15.0) * 100))

        # 5. Education (10% of total score)
        edu_count = len(education) or (1 if sections.get('education') else 0)
        if edu_count >= 1:
            education_score = 10
        elif sections.get('education'):
            education_score = 7
        else:
            education_score = 3
        education_pct = min(100, int((education_score / 10.0) * 100))

        # 6. Formatting & Structure (10% of total score)
        word_count = len(text.split())
        has_summary = bool(sections.get('summary'))
        has_skills_sec = bool(sections.get('skills'))
        formatting_score = 10 if (200 <= word_count <= 1500 and (has_summary or has_skills_sec)) else (7 if word_count >= 100 else 4)
        formatting_pct = min(100, int((formatting_score / 10.0) * 100))

        # 7. Achievements & Measurable Metrics (5% of total score)
        metric_count = len(re.findall(r'(\d+[\%|\+]|\$\d+|\d+\s*(users|clients|teams|projects|x))', text_lower))
        action_count = sum(1 for verb in cls.ACTION_VERBS if re.search(r'\b' + verb + r'\b', text_lower))
        if metric_count >= 2 or action_count >= 4:
            achievement_score = 5
        elif metric_count >= 1 or action_count >= 2:
            achievement_score = 3
        else:
            achievement_score = 1
        achievement_pct = min(100, int((achievement_score / 5.0) * 100))

        # 8. Industry Core Keywords Match (20% of total score)
        core_keywords = ['git', 'sql', 'rest', 'api', 'data', 'cloud', 'testing', 'docker', 'agile', 'system']
        found_kw_count = sum(1 for kw in core_keywords if kw in text_lower)
        keyword_score = min(20, max(4, round((found_kw_count / len(core_keywords)) * 20.0)))
        keyword_pct = min(100, int((keyword_score / 20.0) * 100))

        # Compute Transparent ATS Overall Score (Total = 100)
        ats_score = min(100, max(20, round(
            keyword_score +
            skill_score +
            experience_score +
            project_score +
            education_score +
            formatting_score +
            achievement_score +
            contact_pts
        )))

        breakdown = {
            'keyword_match': keyword_pct,
            'skills': skill_pct,
            'experience': experience_pct,
            'projects': project_pct,
            'education': education_pct,
            'formatting': formatting_pct,
            'achievements': achievement_pct,
            'contact_information': contact_pct,
            'ats_compatibility': ats_score,
        }

        overall_score = ats_score

        strengths = []
        improvements = []

        if skill_count >= 6:
            strengths.append(f"Strong skill representation with {skill_count} relevant proficiencies detected.")
        else:
            improvements.append("Expand your technical skills section with more role-specific technologies.")

        if action_count >= 3:
            strengths.append("Effective use of strong action verbs across accomplishment bullet points.")
        else:
            improvements.append("Begin project and work bullets with strong action verbs (e.g. 'Engineered', 'Optimized').")

        if metric_count >= 1:
            strengths.append("Good inclusion of measurable metrics, percentages, and quantitative results.")
        else:
            improvements.append("Add measurable outcomes (e.g. 'improved efficiency by 25%', 'served 1,000+ users').")

        if has_summary:
            strengths.append("Clear professional summary statement included at the top.")
        else:
            improvements.append("Include a concise 2-3 sentence professional summary at the top.")

        if not has_link:
            improvements.append("Include a link to your GitHub or LinkedIn profile for recruiter verification.")

        # Keywords commonly recommended for tech roles
        common_essential = ['Git', 'Docker', 'REST APIs', 'SQL', 'CI/CD Pipelines', 'System Design']
        missing_keywords = [k for k in common_essential if k not in skills][:4]

        scores = {
            'ats': ats_score,
            'overall': overall_score,
            'breakdown': breakdown,
        }

        return scores, strengths, improvements, missing_keywords

    @classmethod
    def _evaluate_career_alignment(cls, normalized_skills: list, target_career=None) -> dict:
        from career.models import CareerPath, SkillRequirement

        if not target_career:
            target_career = CareerPath.objects.first()

        if not target_career:
            return {
                'target_role': 'Software / AI Engineer',
                'match_percentage': 70,
                'strong_skills': normalized_skills[:5],
                'missing_skills': ['Docker', 'CI/CD Pipelines'],
            }

        reqs = list(SkillRequirement.objects.filter(career_path=target_career).select_related('skill'))
        if not reqs:
            req_skill_names = [s.name for s in target_career.required_skills.all()]
        else:
            req_skill_names = [r.skill.name for r in reqs]

        if not req_skill_names:
            req_skill_names = ['Python', 'SQL', 'Machine Learning', 'Data Structures & Algorithms', 'Git']

        strong = []
        missing = []

        resume_skills_lower = {s.lower() for s in normalized_skills}

        for req_name in req_skill_names:
            if req_name.lower() in resume_skills_lower:
                strong.append(req_name)
            else:
                missing.append(req_name)

        total_req = len(req_skill_names) or 1
        match_pct = round((len(strong) / total_req) * 100)

        return {
            'career_id': target_career.id,
            'target_role': target_career.title,
            'match_percentage': match_pct,
            'strong_skills': strong,
            'missing_skills': missing,
        }

    @classmethod
    def _empty_or_failed_result(cls, filename: str, msg: str) -> dict:
        return {'error': msg}
